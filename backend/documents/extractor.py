from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


class DocumentExtractionError(ValueError):
    """Raised when a supported document cannot produce usable text."""


class UnsupportedDocumentTypeError(DocumentExtractionError):
    """Raised when a document extension is not supported."""


def supported_extension(filename: str) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedDocumentTypeError(
            f"Unsupported file type. Supported extensions: {supported}"
        )
    return extension


def _extract_utf8_text(file_path: Path, file_type: str) -> str:
    try:
        text = file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError as error:
        raise DocumentExtractionError(
            f"{file_type} files must use UTF-8 encoding."
        ) from error

    if not text.strip():
        raise DocumentExtractionError(
            f"No extractable text was found in the {file_type} file."
        )
    return text


def extract_txt_text(file_path: Path) -> str:
    return _extract_utf8_text(file_path, "TXT")


def extract_markdown_text(file_path: Path) -> str:
    """Treat Markdown as UTF-8 plain text; no Markdown parser is needed."""
    return _extract_utf8_text(file_path, "Markdown")


def extract_pdf_text(file_path: Path) -> str:
    try:
        reader = PdfReader(file_path)
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except (PdfReadError, OSError, ValueError) as error:
        raise DocumentExtractionError("The PDF could not be read.") from error

    if not text.strip():
        raise DocumentExtractionError(
            "No extractable text was found. The PDF may be scanned and require OCR."
        )
    return text


def extract_docx_text(file_path: Path) -> str:
    try:
        document = Document(file_path)
    except (PackageNotFoundError, BadZipFile, KeyError, OSError, ValueError) as error:
        raise DocumentExtractionError("The DOCX file could not be read.") from error

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))

    text = "\n".join(parts)
    if not text.strip():
        raise DocumentExtractionError("No extractable text was found in the DOCX file.")
    return text


def extract_text(file_path: Path | str) -> str:
    path = Path(file_path)
    extension = supported_extension(path.name)

    extractors = {
        ".txt": extract_txt_text,
        ".md": extract_markdown_text,
        ".pdf": extract_pdf_text,
        ".docx": extract_docx_text,
    }
    return extractors[extension](path)
